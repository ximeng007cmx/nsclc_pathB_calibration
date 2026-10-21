#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
27_make_tables.py
生成 Table 1–3：
  Table 1  49 签名清单（注册表自动生成，含 PMID/DOI，核验状态另见 refs_verify 输出）
  Table 2  多口径汇总（ALL/W1–W4/三分层，bootstrap CI，summary_all.tsv 直读）
  Table 3  外部复算 + 终点敏感性（external_m1m3_summary / summary_endpoints / rho 直读）
输出：manuscript/tables.xlsx（3 个 sheet）+ manuscript/tables.md
纪律：所有数字直读并 assert（与 06 号修订口径一致）。
"""
import csv, os, json

B = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = os.path.join(B, "results")
W = os.path.join(B, "work")
OUTX = os.path.join(B, "manuscript", "tables.xlsx")
OUTM = os.path.join(B, "manuscript", "tables.md")


def rd(name, base=None, delim=","):
    with open(os.path.join(base or B, name), encoding="utf-8-sig", newline="") as f:
        return [dict(r) for r in csv.DictReader(f, delimiter=delim)]


def F(x):
    try:
        return float(x)
    except Exception:
        return None


def fmt(v, d=1, plus=False):
    v = F(v)
    if v is None:
        return "—"
    return ("%+." + str(d) + "f") % v if plus else ("%." + str(d) + "f") % v


# ============ Table 1：注册表 → 49 行 ============
files = [("1", "signature_registry.csv"), ("2", "signature_registry_expansion.csv"),
         ("3", "signature_registry_expansion_w3.csv"), ("4", "signature_registry_expansion_w4.csv")]
t1 = []
for wave, fn in files:
    for r in rd(fn):
        t1.append({
            "sig_id": r["sig_id"], "wave": wave, "stratum": r["stratum"],
            "journal": r["journal"], "year": r["year"], "pmid": r["pmid"], "doi": r["doi"],
            "n_genes": r["n_genes"], "title": r["title"].strip(),
            "pool_evidence": (r.get("evidence") or "").strip(),
        })
assert len(t1) == 49, len(t1)
by_stratum = {}
for r in t1:
    by_stratum[r["stratum"]] = by_stratum.get(r["stratum"], 0) + 1
assert by_stratum == {"inflammation": 11, "metabolism": 32, "immune": 6}, by_stratum
order = {r["sig_id"]: i for i, r in enumerate(t1)}
t1.sort(key=lambda r: order[r["sig_id"]])

# ============ Table 2：多口径汇总 ============
SUM = rd("summary_all.tsv", R, "\t")
S = {r["set"]: r for r in SUM}
need = ["ALL", "W1", "W2", "W3", "W4", "inflammation", "metabolism", "immune"]
assert all(k in S for k in need), list(S)
assert F(S["ALL"]["median_A"]) is not None and abs(F(S["ALL"]["median_A"]) - (-1.379)) < 0.01
t2 = []
for k in need:
    r = S[k]
    lab = {"ALL": "All signatures", "W1": "Wave 1", "W2": "Wave 2", "W3": "Wave 3", "W4": "Wave 4",
           "inflammation": "Inflammation stratum", "metabolism": "Metabolism stratum",
           "immune": "Immune stratum"}[k]
    t2.append({
        "set": lab, "n": int(r["n"]),
        "median_A": fmt(r["median_A"], 1, True),
        "ci": "[%s, %s]" % (fmt(r["lo"], 1, True), fmt(r["hi"], 1, True)) if r["lo"] else "—",
        "n_valid": int(r["n_ratio_valid"]),
        "median_A_valid": fmt(r["median_A_valid"], 1, True),
        "ci_valid": "[%s, %s]" % (fmt(r["lo_valid"], 1, True), fmt(r["hi_valid"], 1, True)) if r["lo_valid"] else "—",
        "n_gt100": int(r["n_abs_gt100"]), "n_flip": int(r["n_dir_flip"]),
        "fdr": "%d / %d" % (int(r["n_FDR_M1"]), int(r["n_FDR_M3"])),
    })

# ============ Table 3：外部复算 + 终点敏感性 ============
EXT = rd("external_m1m3_summary.tsv", R, "\t")
E = {(r["cohort"], r["stratum"]): r for r in EXT}
EXTRHO = {r["cohort"]: r for r in rd("external_vs_tcga_rho.tsv", R, "\t")}
SUMEP = {(r["endpoint"], r["set"]): r for r in rd("summary_endpoints.tsv", R, "\t")}
strata3 = [("immune", "Immune (6)"), ("inflammation", "Inflammation (11)"), ("metabolism", "Metabolism (31*)")]
cohorts = [("GSE68465", "GSE68465 (n = 439, 235 events)"), ("GSE31210", "GSE31210 (n = 226, 35 events)")]
t3a = []
for coh, clab in cohorts:
    for st, slab in strata3:
        r = E[(coh, st)]
        n_eff = int(r["n"])
        if coh == "GSE68465" and st == "metabolism":
            assert n_eff == 31
        t3a.append({"cohort": clab, "stratum": slab, "n": n_eff,
                    "median_A": fmt(F(r["median_A"]) * 100, 1, True),
                    "n_pos": int(r["n_pos"]), "n_neg": int(r["n_neg"])})
# TCGA 对照列（summary_all，已是百分数）
t3a_tcga = {st: fmt(F(S[st]["median_A"]), 1, True) for st in ("immune", "inflammation", "metabolism")}

t3b = []
for ep, elab in (("OS", "Overall survival (n = 482, 177 events)"),
                 ("PFI", "Progression-free interval (n = 482, 197 events)"),
                 ("DSS", "Disease-specific survival (n = 447, 107 events)")):
    r = SUMEP[(ep, "ALL")]
    sh, gr, fl = int(r["n_pos"]), int(r["n_neg"]), int(r["n_flip"])
    t3b.append({"endpoint": elab, "median_A": fmt(r["median_A"], 1, True),
                "ci": "[%s, %s]" % (fmt(r["lo"], 1, True), fmt(r["hi"], 1, True)),
                "dirs": "%d / %d / %d" % (sh - fl, gr, fl), "fdr": "%d / %d" % (int(r["n_FDR_M1"]), int(r["n_FDR_M3"]))})
for ep in ("OS", "PFI", "DSS"):
    for st in ("immune", "inflammation", "metabolism"):
        v = F(SUMEP[(ep, st)]["median_A"])
        assert v is not None
assert abs(F(EXTRHO["GSE68465"]["rho"]) - (-0.014)) < 0.001
assert abs(F(EXTRHO["GSE31210"]["rho"]) - 0.006) < 0.001

# 跨终点 ρ（与 3.8 一致）
EP = rd("attenuation_endpoints.tsv", R, "\t")


def spearman(x, y):
    def rk(v):
        n = len(v); srt = sorted(range(n), key=lambda i: v[i]); r = [0.0] * n; i = 0
        while i < n:
            j = i
            while j + 1 < n and v[srt[j + 1]] == v[srt[i]]:
                j += 1
            for k in range(i, j + 1):
                r[srt[k]] = (i + j) / 2.0 + 1
            i = j + 1
        return r
    return spearman_pair(rk(x), rk(y))


def spearman_pair(a, b):
    n = len(a); ma = sum(a) / n; mb = sum(b) / n
    num = sum((a[i] - ma) * (b[i] - mb) for i in range(n))
    den = math_sqrt(sum((v - ma) ** 2 for v in a) * sum((v - mb) ** 2 for v in b))
    return num / den if den else float("nan")


def math_sqrt(x):
    return x ** 0.5


A_by = {}
for r in EP:
    A_by.setdefault(r["endpoint"], {})[r["sig_id"]] = F(r["A"])
rho_ep = {}
for e2 in ("PFI", "DSS"):
    ids = [k for k in A_by["OS"] if k in A_by[e2]]
    rho_ep[e2] = spearman([A_by["OS"][k] for k in ids], [A_by[e2][k] for k in ids])
assert abs(rho_ep["PFI"] - 0.418) < 0.002 and abs(rho_ep["DSS"] - 0.639) < 0.002, rho_ep

# ============ xlsx ============
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_font = Font(color="FFFFFF", bold=True, size=10)


def sheet(ws, headers, rows, widths, wrap_cols=()):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=(j in wrap_cols))
    for j, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A2"


ws1 = wb.active; ws1.title = "Table1_signatures"
sheet(ws1,
      ["SIG", "Wave", "Stratum", "First-journal", "Year", "PMID", "DOI", "n_genes", "Title (as screened)", "Pool provenance note"],
      [[r["sig_id"], r["wave"], r["stratum"], r["journal"], r["year"], r["pmid"], r["doi"],
        int(r["n_genes"]), r["title"], r["pool_evidence"][:300] or "—"] for r in t1],
      [8, 7, 14, 20, 7, 12, 30, 9, 70, 40], wrap_cols=(9, 10))

ws2 = wb.create_sheet("Table2_calibers")
sheet(ws2,
      ["Set", "n", "Median A % (all)", "95% CI", "n ratio-valid", "Median A % (valid)", "95% CI (valid)",
       "n |A|>100%", "n flip", "FDR<0.05 M1 / M3"],
      [[r["set"], r["n"], r["median_A"], r["ci"], r["n_valid"], r["median_A_valid"], r["ci_valid"],
        r["n_gt100"], r["n_flip"], r["fdr"]] for r in t2],
      [20, 6, 14, 18, 13, 15, 18, 11, 8, 16])

ws3 = wb.create_sheet("Table3_external_endpoints")
ws3["A1"] = "A. External directional replication (M1→M3, stratum medians of A, %)"
ws3["A1"].font = Font(bold=True)
rows_a = [["Cohort", "Stratum", "n", "Median A %", "A>0", "A<0", "TCGA-LUAD median A %"]]
for r in t3a:
    st_key = r["stratum"].split(" ")[0].lower()
    tcga = t3a_tcga["immune" if "immune" in r["stratum"].lower() else
                     ("inflammation" if "inflamm" in r["stratum"].lower() else "metabolism")]
    rows_a.append([r["cohort"], r["stratum"], r["n"], r["median_A"], r["n_pos"], r["n_neg"], tcga])
sheet_start = 3
for j, h in enumerate(rows_a[0], 1):
    c = ws3.cell(sheet_start, j, h); c.fill = hdr_fill; c.font = hdr_font
for i, row in enumerate(rows_a[1:], sheet_start + 1):
    for j, v in enumerate(row, 1):
        ws3.cell(i, j, v)
r0 = sheet_start + len(rows_a) + 2
ws3.cell(r0, 1, "B. Endpoint sensitivity (all 49 signatures, M1→M3)").font = Font(bold=True)
heads_b = ["Endpoint", "Median A %", "95% CI", "Shrink / Grow / Flip", "FDR<0.05 M1 / M3"]
for j, h in enumerate(heads_b, 1):
    c = ws3.cell(r0 + 1, j, h); c.fill = hdr_fill; c.font = hdr_font
for i, r in enumerate(t3b, r0 + 2):
    for j, v in enumerate(r.items(), 1):
        ws3.cell(i, j, v[1])
rf = r0 + len(t3b) + 3
ws3.cell(rf, 1, "Spearman ρ of per-signature A, TCGA vs GSE68465 = %+.3f (n = 48); vs GSE31210 = %+.3f (n = 48)."
         % (F(EXTRHO["GSE68465"]["rho"]), F(EXTRHO["GSE31210"]["rho"])))
ws3.cell(rf + 1, 1, "Spearman ρ of per-signature A across endpoints: OS vs PFI = %+.3f; OS vs DSS = %+.3f."
         % (rho_ep["PFI"], rho_ep["DSS"]))
ws3.cell(rf + 2, 1, "* external metabolism stratum has 31 signatures (SIG26 has no gene on either array platform).")
for j, wd in enumerate([44, 22, 8, 14, 10, 10, 22], 1):
    ws3.column_dimensions[get_column_letter(j)].width = wd

wb.save(OUTX)
print("->", OUTX)

# ============ markdown（可直接粘贴） ============
md = []
md.append("# Tables（MS_draft_v1 配套）\n")
md.append("数字来源：`results/*.tsv` 直读；勾稽见 `06_amendment_numbers.md`。\n")

md.append("## Table 1. The 49 included signatures（注册表自动生成；PMID/DOI 已双源核验，见 logs/refs_verify_49.log）\n")
md.append("| SIG | Wave | Stratum | Journal | Year | PMID | DOI | Genes |")
md.append("|---|---|---|---|---|---|---|---|")
for r in t1:
    md.append("| %s | %s | %s | %s | %s | %s | %s | %s |" %
              (r["sig_id"], r["wave"], r["stratum"], r["journal"], r["year"], r["pmid"], r["doi"], r["n_genes"]))
md.append("\n完整题名与池来源证据见 `tables.xlsx`（Table1_signatures sheet）。\n")

md.append("## Table 2. Attenuation summaries by set (TCGA-LUAD, OS, M1→M3)\n")
md.append("Median A over all signatures and over the ratio-valid subset (nominally significant M1 coefficient); "
          "CIs from 5,000 cluster bootstraps over signatures (seed 20261020).\n")
md.append("| Set | n | Median A, % | 95% CI | n valid | Median A (valid), % | 95% CI | \\|A\\|>100% | Flip | FDR M1/M3 |")
md.append("|---|---|---|---|---|---|---|---|---|---|")
for r in t2:
    md.append("| %s | %d | %s | %s | %d | %s | %s | %d | %d | %s |" %
              (r["set"], r["n"], r["median_A"], r["ci"], r["n_valid"], r["median_A_valid"], r["ci_valid"],
               r["n_gt100"], r["n_flip"], r["fdr"]))

md.append("\n## Table 3. External directional replication and endpoint sensitivity\n")
md.append("**A. Stratum medians of A (%) by cohort.**\n")
md.append("| Cohort | Stratum | n | Median A, % | A>0 | A<0 | TCGA-LUAD median A, % |")
md.append("|---|---|---|---|---|---|---|")
for r in t3a:
    key = "immune" if "immune" in r["stratum"].lower() else \
          ("inflammation" if "inflamm" in r["stratum"].lower() else "metabolism")
    md.append("| %s | %s | %d | %s | %d | %d | %s |" %
              (r["cohort"], r["stratum"], r["n"], r["median_A"], r["n_pos"], r["n_neg"], t3a_tcga[key]))
md.append("\n**B. Endpoint sensitivity (all 49 signatures).**\n")
md.append("| Endpoint | Median A, % | 95% CI | Shrink / Grow / Flip | FDR M1/M3 |")
md.append("|---|---|---|---|---|")
for r in t3b:
    md.append("| %s | %s | %s | %s | %s |" % (r["endpoint"], r["median_A"], r["ci"], r["dirs"], r["fdr"]))
md.append("\nPer-signature A across cohorts: ρ(TCGA, GSE68465) = %+.3f, ρ(TCGA, GSE31210) = %+.3f (n = 48; SIG26 "
          "unscoreable externally, hence metabolism n = 31*). Across endpoints: ρ(OS, PFI) = %+.3f, ρ(OS, DSS) = %+.3f."
          % (F(EXTRHO["GSE68465"]["rho"]), F(EXTRHO["GSE31210"]["rho"]), rho_ep["PFI"], rho_ep["DSS"]))
md.append("\nDirection counts exclude reversals (3 per endpoint); FDR columns give q<0.05 counts before/after adjustment.\n")

with open(OUTM, "w", encoding="utf-8") as f:
    f.write("\n".join(md))
print("->", OUTM)
print("t1=%d t2=%d t3a=%d t3b=%d" % (len(t1), len(t2), len(t3a), len(t3b)))
